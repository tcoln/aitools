# Author: glt
# Email: guolintan@qq.com
# Created: 2026-07-22

import io
import logging
import os
import subprocess
import tempfile

logger = logging.getLogger(__name__)

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    logger.warning("openpyxl 未安装，Excel 文件解析不可用")

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    logger.warning("python-docx 未安装，Word 文件解析不可用")

try:
    from pypdf import PdfReader
    HAS_PYPDF = True
except ImportError:
    HAS_PYPDF = False
    logger.warning("pypdf 未安装，PDF 文件解析不可用")

XLSX_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
DOCX_EXTENSIONS = {".docx", ".docm", ".dotx", ".dotm", ".wpsx"}
LEGACY_EXTENSIONS = {
    ".xls": "旧版 Excel 格式(.xls)不支持，请另存为 .xlsx 格式",
    ".ppt": "旧版 PPT 格式(.ppt)不支持，请另存为 .pptx 格式",
}
LEGACY_BINARY_EXTENSIONS = {".doc", ".wps", ".et"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".py", ".js", ".ts", ".jsx", ".tsx", ".json", ".csv",
    ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".log", ".html",
    ".css", ".sql", ".sh", ".bash", ".env", ".gitignore", ".java", ".c",
    ".cpp", ".h", ".rs", ".go", ".rb", ".php", ".swift", ".kt", ".scala",
    ".r", ".lua", ".vim", ".conf",
}

PDF_EXTENSIONS = {".pdf"}


# 根据文件扩展名解析文件内容，支持xlsx、docx、文本、pdf、旧版doc/wps/et等格式
def parse_file(filename: str, content: bytes) -> str:
    ext = filename.lower()
    for e in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        if ext.endswith(e):
            return _parse_xlsx(content)
    for e in (".docx", ".docm", ".dotx", ".dotm", ".wpsx"):
        if ext.endswith(e):
            return _parse_docx(content)
    for e in PDF_EXTENSIONS:
        if ext.endswith(e):
            return _parse_pdf(content)
    for e in LEGACY_BINARY_EXTENSIONS:
        if ext.endswith(e):
            return _parse_legacy_binary(filename, content)
    for e, msg in LEGACY_EXTENSIONS.items():
        if ext.endswith(e):
            return f"[{msg}]"
    for e in TEXT_EXTENSIONS:
        if ext.endswith(e):
            return _parse_text(content)
    return f"[不支持解析此文件格式: {filename}]"


# 解析Excel(.xlsx)文件内容
def _parse_xlsx(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "[Excel 解析失败: openpyxl 未安装，请运行 pip install openpyxl]"
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                rows.append(" | ".join(cells))
            if rows:
                parts.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
        wb.close()
        return "\n\n".join(parts) if parts else "[空表格]"
    except Exception as e:
        logger.error(f"Failed to parse xlsx: {e}")
        return f"[Excel 解析失败: {e}]"


# 解析Word(.docx)文件内容
def _parse_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        return "[Word 解析失败: python-docx 未安装，请运行 pip install python-docx]"
    try:
        doc = Document(io.BytesIO(content))
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text for cell in row.cells]
                tables_text.append(" | ".join(cells))
        result = "\n".join(paragraphs)
        if tables_text:
            result += "\n\n--- 表格内容 ---\n" + "\n".join(tables_text)
        return result if result.strip() else "[空文档]"
    except Exception as e:
        logger.error(f"Failed to parse docx: {e}")
        return f"[Word 文档解析失败: {e}]"


# 解析文本文件内容，支持UTF-8和GBK编码
def _parse_text(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        try:
            return content.decode("gbk")
        except Exception:
            return "[无法解码文件内容]"


# 解析PDF文件内容
def _parse_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[PDF 解析失败: pypdf 未安装，请运行 pip install pypdf]"
    try:
        reader = PdfReader(io.BytesIO(content))
        parts = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            if text and text.strip():
                parts.append(f"--- 第 {i} 页 ---\n{text.strip()}")
        if not parts:
            return "[PDF 内容为空或无法提取文本]"
        return "\n\n".join(parts)
    except Exception as e:
        logger.error(f"Failed to parse pdf: {e}")
        return f"[PDF 解析失败: {e}]"


# 解析旧版二进制格式文件(.doc/.wps/.et)，通过 LibreOffice 转换
def _parse_legacy_binary(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    format_names = {".doc": "旧版 Word", ".wps": "WPS 文字", ".et": "WPS 表格"}

    try:
        result = _parse_via_libreoffice(content, ext)
        if result is not None:
            return result
    except Exception as e:
        logger.warning(f"LibreOffice 转换失败: {e}")

    friendly_name = format_names.get(ext, ext)
    return f"[{friendly_name}格式({ext})解析失败，请确保服务器已安装 LibreOffice，或另存为 docx/xlsx 格式]"


# 使用 LibreOffice headless 模式将文档转换为纯文本
def _parse_via_libreoffice(content: bytes, ext: str) -> str | None:
    try:
        subprocess.run(
            ["libreoffice", "--version"],
            capture_output=True, timeout=5,
        )
    except (FileNotFoundError, subprocess.TimeoutExpired):
        logger.warning("LibreOffice 未安装或不可用")
        return None

    tmpdir = tempfile.mkdtemp(prefix="aitools_parse_")
    src_path = os.path.join(tmpdir, f"input{ext}")

    try:
        with open(src_path, "wb") as f:
            f.write(content)

        proc = subprocess.run(
            [
                "libreoffice", "--headless", "--convert-to", "txt:Text",
                "--outdir", tmpdir, src_path,
            ],
            capture_output=True, timeout=60,
        )

        if proc.returncode != 0:
            stderr = proc.stderr.decode("utf-8", errors="replace")[:500]
            logger.error(f"LibreOffice 转换失败: {stderr}")
            return None

        base_name = os.path.splitext(os.path.basename(src_path))[0]
        txt_path = os.path.join(tmpdir, f"{base_name}.txt")

        if not os.path.exists(txt_path):
            logger.error(f"LibreOffice 输出文件不存在: {txt_path}")
            return None

        with open(txt_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()

        if not text.strip():
            return "[文档内容为空]"

        return text
    except subprocess.TimeoutExpired:
        logger.error("LibreOffice 转换超时")
        return None
    except Exception as e:
        logger.error(f"LibreOffice 转换异常: {e}")
        return None
    finally:
        try:
            for f in os.listdir(tmpdir):
                os.unlink(os.path.join(tmpdir, f))
            os.rmdir(tmpdir)
        except Exception:
            pass